from __future__ import annotations

import re
from collections.abc import Iterator
from pathlib import Path

from codecite.parse.base import Block


def parse_docx(path: Path) -> Iterator[Block]:
    import docx

    d = docx.Document(str(path))
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = p.style.name or ""
        m = re.match(r"Heading (\d)", style)
        if m:
            yield Block("heading", t, level=int(m.group(1)))
        elif style.lower().startswith("list"):
            yield Block("list-item", t)
        else:
            yield Block("paragraph", t)
    for tbl in d.tables:
        rows = []
        for i, row in enumerate(tbl.rows):
            cells = [c.text.strip() for c in row.cells]
            rows.append("| " + " | ".join(cells) + " |")
            if i == 0:
                rows.append("|" + "---|" * len(cells))
        yield Block("table", "\n".join(rows))
